"""Formal data-task runner for the authorized 2026-Q2 base replacement.

It reads only the caller-provided original workbook, streams the permitted seven
source fields to an existing server-side PostgreSQL container, and fails closed.
`--environment formal` is intentionally unavailable without a successful staging
report supplied on the command line.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shlex
import subprocess
import sys
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

EXPECTED_SHA256 = "6f29945a14471fc9be4dfff34d40dfeb045239179e92dd640e719cfdce16bb73"
EXPECTED_HEADERS = ["序号", "对账时间点", "账套", "区域", "对账负责人", "客户名称", "公司应收"]
EXPECTED_ACCOUNTS = {"国药控股": 79, "英科": 552, "明生医疗": 78, "万和": 31, "生一": 39, "江苏鑫之阳": 2, "道壹检测": 69}
EXPECTED_REGIONS = {"苏州":171,"南京":135,"南通":131,"无锡":105,"盐城":43,"常州":40,"扬州":39,"徐州":38,"淮安":34,"血站":32,"镇江":25,"泰州":23,"宿迁":16,"连云港":12,"三方/淮安":3,"公司":2,"三方/镇江":1}
EXPECTED_RECEIVABLE = Decimal("518644096.68")


def fail(message: str) -> None:
    raise RuntimeError(message)


def sql_string(value: object | None) -> str:
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def parse_source(path: Path) -> tuple[str, list[dict[str, object | None]]]:
    if not path.is_file():
        fail(f"Source workbook does not exist: {path}")
    source_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    if source_hash != EXPECTED_SHA256:
        fail("SOURCE_SHA256_MISMATCH")
    workbook = load_workbook(path, read_only=True, data_only=False)
    if len(workbook.worksheets) != 1:
        fail("SOURCE_WORKSHEET_COUNT_MISMATCH")
    sheet = workbook.worksheets[0]
    headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
    if headers[:7] != EXPECTED_HEADERS:
        fail("SOURCE_ALLOWED_COLUMNS_MISMATCH")
    rows: list[dict[str, object | None]] = []
    for raw in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in raw):
            continue
        values = raw[:7]
        sequence, timepoint, account, region, owner, customer, receivable = values
        if sequence is None or not str(account).strip() or not str(region).strip() or not str(customer).strip():
            fail("SOURCE_REQUIRED_VALUE_MISSING")
        amount = None if receivable is None else Decimal(str(receivable)).quantize(Decimal("0.01"))
        key = "q2:" + hashlib.sha256(f"{source_hash}|{sequence}".encode()).hexdigest()[:24]
        rows.append({"sequence": str(sequence), "timepoint": None if timepoint is None else str(timepoint), "account": str(account).strip(), "region": str(region).strip(), "owner": None if owner is None else str(owner).strip(), "customer": str(customer).strip(), "receivable": amount, "source_row_key": key})
    accounts = Counter(str(row["account"]) for row in rows)
    regions = Counter(str(row["region"]) for row in rows)
    amounts = [row["receivable"] for row in rows]
    blank_owners = sum(row["owner"] in (None, "") for row in rows)
    zero_owners = sum(row["owner"] == "0" for row in rows)
    sequences = {row["sequence"] for row in rows}
    if not (len(rows) == 850 and accounts == EXPECTED_ACCOUNTS and regions == EXPECTED_REGIONS and amounts.count(None) == 1 and sum(value for value in amounts if value is not None) == EXPECTED_RECEIVABLE and blank_owners == 6 and zero_owners == 16 and "710" not in sequences):
        fail("SOURCE_PRE_VALIDATION_FAILED")
    for sequence in ("275", "514", "712", "722"):
        if sequence not in sequences:
            fail(f"SOURCE_SPECIAL_ROW_MISSING_{sequence}")
    by_sequence = {str(row["sequence"]): row for row in rows}
    if not (by_sequence["275"]["timepoint"] == "26.7" and by_sequence["275"]["account"] == "英科" and by_sequence["275"]["region"] == "泰州" and by_sequence["275"]["owner"] == "董永平" and by_sequence["275"]["customer"] == "南京医药集团股份有限公司" and by_sequence["275"]["receivable"] == Decimal("658694.23")):
        fail("SOURCE_SPECIAL_ROW_275_MISMATCH")
    if not (by_sequence["514"]["account"] == "英科" and by_sequence["514"]["region"] == "南京" and by_sequence["514"]["owner"] == "刘丽婷" and by_sequence["514"]["customer"] == "东南大学附属中大医院（江北院区）" and by_sequence["514"]["receivable"] is None):
        fail("SOURCE_SPECIAL_ROW_514_MISMATCH")
    for sequence in ("712", "722"):
        row = by_sequence[sequence]
        if not (row["account"] == "万和" and row["region"] == "苏州" and row["owner"] == "贾涛" and row["customer"] == "苏州市相城区漕湖人民医院" and row["receivable"] == Decimal("231488.60")):
            fail(f"SOURCE_SPECIAL_ROW_{sequence}_MISMATCH")
    return source_hash, rows


def run_ssh(ssh_target: str, ssh_port: int, remote_command: str, stdin: str | None = None) -> str:
    result = subprocess.run(["ssh", "-p", str(ssh_port), ssh_target, remote_command], input=stdin, text=True, encoding="utf-8", stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    if result.returncode:
        fail(f"REMOTE_COMMAND_FAILED: {result.stderr.strip()}")
    return result.stdout


def psql_command(container: str, db_user: str) -> str:
    # No credentials are read or emitted locally: psql runs inside the existing DB container.
    return f"set -eu; docker exec -i {shlex.quote(container)} sh -c 'psql -v ON_ERROR_STOP=1 -U {shlex.quote(db_user)} -d \"$POSTGRES_DB\"'"


def backup_command(container: str, db_user: str, environment: str) -> tuple[str, str]:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    if environment == "formal":
        path = f"/backups/formal-replace-2026-q2-{stamp}.dump"
        command = f"set -eu; docker exec {shlex.quote(container)} sh -c 'pg_dump -U {shlex.quote(db_user)} -d \"$POSTGRES_DB\" -Fc -f {shlex.quote(path)}; sha256sum {shlex.quote(path)}'"
    else:
        path = f"/tmp/staging-replace-2026-q2-{stamp}.dump"
        command = f"set -eu; docker exec {shlex.quote(container)} sh -c 'pg_dump -U {shlex.quote(db_user)} -d \"$POSTGRES_DB\" -Fc' > {shlex.quote(path)}; sha256sum {shlex.quote(path)}"
    return path, command


def build_sql(source_hash: str, rows: list[dict[str, object | None]], environment: str) -> str:
    values = []
    for row in rows:
        payload = json.dumps({"allowed_source_columns": EXPECTED_HEADERS, "row": [row[key] if key != "receivable" else (None if row[key] is None else format(row[key], ".2f")) for key in ("sequence", "timepoint", "account", "region", "owner", "customer", "receivable")]}, ensure_ascii=False, separators=(",", ":"))
        values.append("(" + ",".join([sql_string(row["sequence"]), sql_string(row["timepoint"]), sql_string(row["account"]), sql_string(row["region"]), sql_string(row["owner"]), sql_string(row["customer"]), "NULL" if row["receivable"] is None else str(row["receivable"]), sql_string(row["source_row_key"]), sql_string(payload) + "::jsonb"]) + ")")
    task_key = f"formal-data-task:replace-2026-q2-base-reconciliation:{source_hash}"
    return f"""
\\set ON_ERROR_STOP on
BEGIN;
SELECT pg_advisory_xact_lock(hashtext({sql_string(task_key)}));
CREATE TEMP TABLE task_source (sequence text NOT NULL, timepoint text, account_set text NOT NULL, region text NOT NULL, owner_raw text, customer_name text NOT NULL, company_receivable numeric(18,2), source_row_key text NOT NULL, source_payload jsonb NOT NULL) ON COMMIT DROP;
INSERT INTO task_source VALUES\n{',\n'.join(values)};
DO $$ BEGIN
  IF (SELECT count(*) FROM task_source) <> 850 THEN RAISE EXCEPTION 'PRE_SOURCE_ROWS'; END IF;
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q1') <> 752 THEN RAISE EXCEPTION 'PRE_Q1'; END IF;
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2') <> 851 THEN RAISE EXCEPTION 'PRE_Q2'; END IF;
  IF (SELECT count(*) FROM recon.historical_settlement_snapshots) <> 108 THEN RAISE EXCEPTION 'PRE_HISTORICAL'; END IF;
  IF (SELECT count(*) FROM recon.schema_migrations) <> 9 THEN RAISE EXCEPTION 'PRE_MIGRATIONS'; END IF;
  IF EXISTS (SELECT 1 FROM task_source s LEFT JOIN recon.account_sets a ON a.name=s.account_set WHERE a.id IS NULL) THEN RAISE EXCEPTION 'PRE_ACCOUNT_SET_DICTIONARY'; END IF;
END $$;
INSERT INTO recon.regions (code, name)
SELECT additions.name, additions.name
FROM (VALUES ('三方/淮安'), ('三方/镇江')) AS additions(name)
LEFT JOIN recon.regions existing ON existing.name = additions.name
WHERE existing.id IS NULL;
DO $$ BEGIN
  IF EXISTS (SELECT 1 FROM task_source s LEFT JOIN recon.regions g ON g.name=s.region WHERE g.id IS NULL) THEN RAISE EXCEPTION 'PRE_REGION_DICTIONARY'; END IF;
  IF (SELECT count(*) FROM recon.regions WHERE name IN ('三方/淮安','三方/镇江')) <> 2 THEN RAISE EXCEPTION 'PRE_REGION_EXACT'; END IF;
  IF (SELECT count(*) FROM task_source WHERE region='三方') <> 0 THEN RAISE EXCEPTION 'PRE_SOURCE_STANDALONE_THIRD_PARTY'; END IF;
END $$;
CREATE TEMP TABLE old_q2_customers AS SELECT DISTINCT r.customer_id AS id FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2';
DELETE FROM recon.followup_events e USING recon.followup_items f, recon.reconciliations r, recon.quarters q WHERE e.followup_item_id=f.id AND f.reconciliation_id=r.id AND r.quarter_id=q.id AND q.code='2026-Q2';
DELETE FROM recon.followup_items f USING recon.reconciliations r, recon.quarters q WHERE f.reconciliation_id=r.id AND r.quarter_id=q.id AND q.code='2026-Q2';
DELETE FROM recon.difference_items d USING recon.reconciliations r, recon.quarters q WHERE d.reconciliation_id=r.id AND r.quarter_id=q.id AND q.code='2026-Q2';
DELETE FROM recon.material_status m USING recon.reconciliations r, recon.quarters q WHERE m.reconciliation_id=r.id AND r.quarter_id=q.id AND q.code='2026-Q2';
DELETE FROM recon.material_status m USING recon.quarters q WHERE m.quarter_id=q.id AND q.code='2026-Q2';
DELETE FROM recon.reconciliations r USING recon.quarters q WHERE r.quarter_id=q.id AND q.code='2026-Q2';
DELETE FROM recon.customers c USING old_q2_customers o WHERE c.id=o.id AND NOT EXISTS (SELECT 1 FROM recon.reconciliations r WHERE r.customer_id=c.id) AND NOT EXISTS (SELECT 1 FROM recon.ledger_invoices l WHERE l.customer_id=c.id);
CREATE TEMP TABLE task_customers (customer_id uuid PRIMARY KEY, account_set text NOT NULL, region text NOT NULL, customer_name text NOT NULL, UNIQUE(account_set, region, customer_name)) ON COMMIT DROP;
INSERT INTO task_customers SELECT gen_random_uuid(), account_set, region, customer_name FROM (SELECT DISTINCT account_set, region, customer_name FROM task_source) s;
INSERT INTO recon.customers (id, external_code, name, region_id) SELECT c.customer_id, NULL, c.customer_name, g.id FROM task_customers c JOIN recon.regions g ON g.name=c.region;
INSERT INTO recon.reconciliations (legacy_id, quarter_id, account_set_id, customer_id, owner_id, company_receivable, customer_book_amount, reconciliation_difference, reconciliation_status, source_row_key, source_payload, owner_name)
SELECT s.sequence, q.id, a.id, c.customer_id, NULL, s.company_receivable, NULL, NULL, NULL, s.source_row_key, s.source_payload, CASE WHEN btrim(coalesce(s.owner_raw,'')) IN ('','0') THEN NULL ELSE btrim(s.owner_raw) END FROM task_source s JOIN recon.quarters q ON q.code='2026-Q2' JOIN recon.account_sets a ON a.name=s.account_set JOIN task_customers c ON c.account_set=s.account_set AND c.region=s.region AND c.customer_name=s.customer_name;
INSERT INTO recon.migration_manifests (batch_key, quarter_code, source_sha256, source_record_count, migrated_record_count, status, error_details, completed_at) VALUES ({sql_string(task_key)}, '2026-Q2', {sql_string(source_hash)}, 850, 850, 'migrated', '[]'::jsonb, now());
DO $$ BEGIN
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2') <> 850 THEN RAISE EXCEPTION 'POST_Q2'; END IF;
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q1') <> 752 THEN RAISE EXCEPTION 'POST_Q1'; END IF;
  IF (SELECT count(*) FROM recon.historical_settlement_snapshots) <> 108 THEN RAISE EXCEPTION 'POST_HISTORICAL'; END IF;
  IF (SELECT coalesce(sum(r.company_receivable),0) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2') <> 518644096.68 THEN RAISE EXCEPTION 'POST_SUM'; END IF;
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2' AND r.company_receivable IS NULL) <> 1 THEN RAISE EXCEPTION 'POST_NULL'; END IF;
  IF (SELECT count(r.customer_book_amount) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2') <> 0 THEN RAISE EXCEPTION 'POST_CUSTOMER_BOOK'; END IF;
  IF (SELECT count(*) FROM recon.difference_items d JOIN recon.reconciliations r ON r.id=d.reconciliation_id JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2') <> 0 THEN RAISE EXCEPTION 'POST_DIFFERENCES'; END IF;
  IF (SELECT count(*) FROM recon.followup_items f JOIN recon.reconciliations r ON r.id=f.reconciliation_id JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2') <> 0 THEN RAISE EXCEPTION 'POST_FOLLOWUPS'; END IF;
  IF (SELECT count(*) FROM recon.material_status m JOIN recon.quarters q ON q.id=m.quarter_id WHERE q.code='2026-Q2') <> 0 THEN RAISE EXCEPTION 'POST_MATERIALS'; END IF;
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2' AND r.owner_name IS NULL) <> 22 THEN RAISE EXCEPTION 'POST_OWNERS'; END IF;
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2' AND r.legacy_id IN ('712','722')) <> 2 THEN RAISE EXCEPTION 'POST_DUPLICATES'; END IF;
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id JOIN recon.customers c ON c.id=r.customer_id JOIN recon.regions g ON g.id=c.region_id WHERE q.code='2026-Q2' AND g.name='三方/淮安') <> 3 THEN RAISE EXCEPTION 'POST_REGION_THIRD_PARTY_HUAIAN'; END IF;
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id JOIN recon.customers c ON c.id=r.customer_id JOIN recon.regions g ON g.id=c.region_id WHERE q.code='2026-Q2' AND g.name='三方/镇江') <> 1 THEN RAISE EXCEPTION 'POST_REGION_THIRD_PARTY_ZHENJIANG'; END IF;
  IF (SELECT count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id JOIN recon.customers c ON c.id=r.customer_id JOIN recon.regions g ON g.id=c.region_id WHERE q.code='2026-Q2' AND g.name='三方') <> 0 THEN RAISE EXCEPTION 'POST_REGION_STANDALONE_THIRD_PARTY'; END IF;
END $$;
SELECT 'DATA_TASK_REPORT|ENVIRONMENT|{environment}';
SELECT 'DATA_TASK_REPORT|Q2|' || count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q2';
SELECT 'DATA_TASK_REPORT|Q1|' || count(*) FROM recon.reconciliations r JOIN recon.quarters q ON q.id=r.quarter_id WHERE q.code='2026-Q1';
SELECT 'DATA_TASK_REPORT|HISTORICAL|' || count(*) FROM recon.historical_settlement_snapshots;
COMMIT;
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--environment", choices=("staging", "formal"), required=True)
    parser.add_argument("--ssh-target", required=True)
    parser.add_argument("--ssh-port", type=int, default=22)
    parser.add_argument("--container", required=True)
    parser.add_argument("--db-user", required=True)
    parser.add_argument("--staging-report", type=Path)
    args = parser.parse_args()
    if args.environment == "formal":
        if not args.staging_report or not args.staging_report.is_file() or '"status": "PASS"' not in args.staging_report.read_text(encoding="utf-8"):
            fail("FORMAL_REQUIRES_SUCCESSFUL_STAGING_REPORT")
    source_hash, rows = parse_source(args.source)
    backup_path, backup = backup_command(args.container, args.db_user, args.environment)
    backup_out = run_ssh(args.ssh_target, args.ssh_port, backup)
    backup_sha = backup_out.strip().split()[0]
    output = run_ssh(args.ssh_target, args.ssh_port, psql_command(args.container, args.db_user), build_sql(source_hash, rows, args.environment))
    report = {"status": "PASS", "environment": args.environment, "source": str(args.source), "source_sha256": source_hash, "backup_path": backup_path, "backup_sha256": backup_sha, "backup_time": datetime.now(timezone.utc).isoformat(), "psql_output": [line for line in output.splitlines() if "DATA_TASK_REPORT" in line]}
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"status": "FAIL", "error": str(error)}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)
