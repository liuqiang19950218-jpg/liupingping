import hashlib
import json
import sys
import zipfile
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

path = Path(sys.argv[1])
expected_hash = "6f29945a14471fc9be4dfff34d40dfeb045239179e92dd640e719cfdce16bb73"
actual_hash = hashlib.sha256(path.read_bytes()).hexdigest()
wb = load_workbook(path, read_only=True, data_only=False)
if len(wb.worksheets) != 1:
    raise SystemExit(f"Expected exactly one worksheet, found {len(wb.worksheets)}")
ws = wb.worksheets[0]
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
expected_headers = ["序号", "对账时间点", "账套", "区域", "对账负责人", "客户名称", "公司应收"]
if headers[:7] != expected_headers:
    raise SystemExit(json.dumps({"header_error": {"actual_first_7": headers[:7], "expected": expected_headers}}, ensure_ascii=False))
rows = []
for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
    if any(value is not None for value in row):
        rows.append((row_num, row[:7]))

accounts = Counter(row[2] for _, row in rows)
regions = Counter(row[3] for _, row in rows)
amounts = [row[6] for _, row in rows]
non_null = [Decimal(str(value)) for value in amounts if value is not None]
with zipfile.ZipFile(path) as zf:
    root = ET.fromstring(zf.read("xl/worksheets/sheet1.xml"))
ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
xml_receivables = []
for cell in root.findall(".//x:c", ns):
    ref = cell.attrib.get("r", "")
    if ref.startswith("G") and ref != "G1":
        value = cell.find("x:v", ns)
        if value is not None and value.text is not None:
            xml_receivables.append(Decimal(value.text))
xml_receivable_cents = [value.quantize(Decimal("0.01")) for value in xml_receivables]
owners = [row[4] for _, row in rows]
blank_owners = sum(value is None or (isinstance(value, str) and value.strip() == "") for value in owners)
zero_owners = sum((isinstance(value, (int, float)) and not isinstance(value, bool) and value == 0) for value in owners)
by_sequence = {row[0]: row for _, row in rows}
special = {str(sequence): list(by_sequence.get(sequence, ())) for sequence in [275, 514, 712, 722]}
expected_accounts = {"国药控股": 79, "英科": 552, "明生医疗": 78, "万和": 31, "生一": 39, "江苏鑫之阳": 2, "道壹检测": 69}
expected_regions = {"苏州":171,"南京":135,"南通":131,"无锡":105,"盐城":43,"常州":40,"扬州":39,"徐州":38,"淮安":34,"血站":32,"镇江":25,"泰州":23,"宿迁":16,"连云港":12,"三方/淮安":3,"公司":2,"三方/镇江":1}
report = {
    "path": str(path), "bytes": path.stat().st_size, "sha256": actual_hash,
    "sha256_matches_expected": actual_hash == expected_hash, "sheet": ws.title,
    "max_row": ws.max_row, "nonempty_source_rows": len(rows), "source_headers": headers,
    "allowed_import_columns": headers[:7], "allowed_import_columns_match": headers[:7] == expected_headers,
    "accounts": dict(sorted(accounts.items())), "accounts_match": accounts == expected_accounts,
    "regions": dict(sorted(regions.items())), "regions_match": regions == expected_regions,
    "receivable_non_null": len(non_null), "receivable_null": len(amounts) - len(non_null),
    "receivable_sum_from_raw_xlsx_values": str(sum(xml_receivables)),
    "receivable_sum_at_currency_precision": str(sum(xml_receivable_cents)),
    "receivable_sum_matches_expected": sum(xml_receivable_cents) == Decimal("518644096.68"),
    "blank_owners": blank_owners, "numeric_zero_owners": zero_owners,
    "missing_sequence_710": 710 not in by_sequence, "special_rows": special,
}
print(json.dumps(report, ensure_ascii=False, indent=2, default=str))
